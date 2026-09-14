from __future__ import annotations

from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "홈페이지_전체_운영_매뉴얼.docx"

FONT = "Malgun Gothic"
MONO = "Consolas"
BLACK = "000000"
BODY = "222222"
MUTED = "666666"
BLUE = "1F4E78"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F9FC"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"
RED = "A61B1B"


def set_run_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.35):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    r_pr.extend([r_fonts, color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


class ManualBuilder:
    def __init__(self):
        self.doc = Document()
        self._configure_document()

    def _configure_document(self):
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.68)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)
        section.header_distance = Inches(0.28)
        section.footer_distance = Inches(0.28)
        section.different_first_page_header_footer = True

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = FONT
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor.from_string(BODY)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.35
        normal.paragraph_format.widow_control = True

        title = styles["Title"]
        title.font.name = FONT
        title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        title.font.size = Pt(28)
        title.font.bold = True
        title.font.color.rgb = RGBColor.from_string(BLACK)
        title.paragraph_format.space_after = Pt(14)
        title_p_pr = title._element.get_or_add_pPr()
        title_border = title_p_pr.find(qn("w:pBdr"))
        if title_border is not None:
            title_p_pr.remove(title_border)

        heading_specs = {
            "Heading 1": (20, 22, 10),
            "Heading 2": (15, 15, 7),
            "Heading 3": (12, 10, 5),
        }
        for name, (size, before, after) in heading_specs.items():
            style = styles[name]
            style.font.name = FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(BLACK)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True

        if "Code Block" not in [s.name for s in styles]:
            code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
            code_style.font.name = MONO
            code_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            code_style.font.size = Pt(8.7)
            code_style.font.color.rgb = RGBColor.from_string(BODY)
            code_style.paragraph_format.left_indent = Inches(0.25)
            code_style.paragraph_format.right_indent = Inches(0.15)
            code_style.paragraph_format.space_before = Pt(3)
            code_style.paragraph_format.space_after = Pt(6)
            code_style.paragraph_format.line_spacing = 1.15

        if "Small Note" not in [s.name for s in styles]:
            note_style = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
            note_style.font.name = FONT
            note_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            note_style.font.size = Pt(9)
            note_style.font.color.rgb = RGBColor.from_string(MUTED)
            note_style.paragraph_format.space_after = Pt(5)
            note_style.paragraph_format.line_spacing = 1.25

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_paragraph_spacing(hp, after=0, line=1.0)
        set_run_font(hp.add_run("바둑타고 한국사 여행  홈페이지 전체 운영 매뉴얼"), size=8.5, color=MUTED)
        add_page_number(section.footer.paragraphs[0])

        settings = self.doc.settings.element
        compatibility = settings.find(qn("w:compat"))
        if compatibility is not None:
            for setting in compatibility:
                if setting.get(qn("w:name")) == "compatibilityMode":
                    setting.set(qn("w:val"), "15")
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)

    def add_body(self, text="", bold_lead=None, style=None, align=None):
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        if bold_lead and text.startswith(bold_lead):
            set_run_font(p.add_run(bold_lead), bold=True)
            set_run_font(p.add_run(text[len(bold_lead):]))
        else:
            set_run_font(p.add_run(text))
        set_paragraph_spacing(p)
        return p

    def add_note(self, lead, text, danger=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.3
        set_run_font(p.add_run(f"{lead}  "), bold=True, color=RED if danger else BLUE)
        set_run_font(p.add_run(text))
        return p

    def add_bullets(self, items, level=0):
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
            p.paragraph_format.left_indent = Inches(0.28 + level * 0.2)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.25
            set_run_font(p.add_run(item))

    def add_steps(self, items):
        for number, item in enumerate(items, start=1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.32)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.3
            set_run_font(p.add_run(f"{number}.  "), bold=True, color=BLUE)
            set_run_font(p.add_run(item))

    def add_checklist(self, items):
        for item in items:
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            p.paragraph_format.space_after = Pt(3)
            set_run_font(p.add_run("[ ] "), bold=True, color=BLUE)
            set_run_font(p.add_run(item))

    def add_code(self, lines):
        text = "\n".join(lines) if isinstance(lines, (list, tuple)) else lines
        p = self.doc.add_paragraph(style="Code Block")
        for idx, line in enumerate(text.splitlines()):
            if idx:
                p.add_run().add_break()
            set_run_font(p.add_run(line), name=MONO, size=8.7)
        return p

    def add_table(self, headers, rows, widths=None, font_size=9.1):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)
        set_repeat_table_header(table.rows[0])
        set_row_cant_split(table.rows[0])
        for idx, header in enumerate(headers):
            cell = table.rows[0].cells[idx]
            cell.text = ""
            set_cell_shading(cell, BLUE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, after=0, line=1.1)
            set_run_font(p.add_run(header), size=9, bold=True, color=WHITE)
            if widths:
                set_cell_width(cell, widths[idx])
        for r_idx, row_data in enumerate(rows):
            cells = table.add_row().cells
            set_row_cant_split(table.rows[-1])
            for c_idx, value in enumerate(row_data):
                cell = cells[c_idx]
                cell.text = ""
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                if r_idx % 2 == 1:
                    set_cell_shading(cell, PALE_BLUE)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                set_paragraph_spacing(p, after=0, line=1.2)
                set_run_font(p.add_run(str(value)), size=font_size)
                if widths:
                    set_cell_width(cell, widths[c_idx])
        after = self.doc.add_paragraph()
        set_paragraph_spacing(after, after=2, line=1.0)
        return table

    def major(self, title, first=False):
        p = self.doc.add_heading(title, level=1)
        p.paragraph_format.keep_with_next = True
        if not first:
            p.paragraph_format.space_before = Pt(18)

    def add_url(self, label, url):
        p = self.doc.add_paragraph()
        set_run_font(p.add_run(f"{label}: "), bold=True)
        add_hyperlink(p, url, url)
        set_paragraph_spacing(p)


def build_manual():
    m = ManualBuilder()
    doc = m.doc

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run("OPERATIONS MANUAL"), size=10, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    set_run_font(title.add_run("바둑타고 한국사 여행\n홈페이지 전체 운영 매뉴얼"), size=28, bold=True, color=BLACK)
    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, after=18, line=1.3)
    set_run_font(subtitle.add_run("콘텐츠 등록부터 회원 지원 결제 점검 서버 배포까지"), size=14, color=MUTED)
    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    set_table_borders(meta)
    cover_rows = [
        ("문서 버전", "1.1"),
        ("운영 기준일", "2026년 9월 10일"),
        ("대상", "운영자 관리자 콘텐츠 담당자 고객지원 담당자 서버 담당자"),
        ("서비스 주소", "https://handol-edu.com"),
    ]
    for ridx, (a, b) in enumerate(cover_rows):
        for cidx, value in enumerate((a, b)):
            cell = meta.rows[ridx].cells[cidx]
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=140, bottom=140)
            set_cell_width(cell, 1.35 if cidx == 0 else 4.9)
            if cidx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            pp = cell.paragraphs[0]
            set_paragraph_spacing(pp, after=0, line=1.2)
            set_run_font(pp.add_run(value), size=9.5, bold=cidx == 0)
    m.add_note("사용 원칙", "일상 운영자는 1장부터 11장까지, 서버 담당자는 12장부터 15장까지 함께 확인합니다. 비밀키와 개인정보는 이 문서나 업무 메신저에 기록하지 않습니다.")

    doc.add_page_break()
    m.major("문서 사용 안내", first=True)
    m.add_body("이 문서는 홈페이지를 실제로 운영할 때 필요한 화면 사용법, 콘텐츠 게시 순서, 고객 응대, 결제 대사, 메일과 소셜 로그인 점검, 서버 배포와 장애 대응을 한 권으로 정리한 절차서입니다. 화면에서 수행하는 업무와 서버에서 수행하는 업무를 분리했으며, 각 기능을 언제 사용하고 어떤 결과를 확인해야 하는지 설명합니다.")
    m.add_note("가장 중요한 결론", "소셜 로그인, 게시판 API, 강의 API는 운영에 연결되어 있습니다. 토스페이먼츠만 테스트 모드이므로 실제 결제를 받기 전에는 라이브 키 전환과 결제 승인, 취소, 환불 검증을 별도로 완료해야 합니다.", danger=True)

    doc.add_heading("목차", level=2)
    m.add_table(
        ["1장부터 8장", "9장부터 15장"],
        [
            ["1. 서비스 현황과 운영 전 주의사항", "9. 구독 결제 대사 환불 운영"],
            ["2. 운영 역할과 권한", "10. 학생 보호자 지도자 기관 기능"],
            ["3. 로그인 계정 보호자 연결", "11. 알림 이메일 소셜 로그인 운영"],
            ["4. 매일 해야 하는 운영 업무", "12. 서버 상태 점검과 로그 확인"],
            ["5. 홈페이지와 시대별 학습 콘텐츠 운영", "13. 배포 변경 승인 롤백"],
            ["6. 강의 콘텐츠 관리", "14. 보안 개인정보 백업 관리"],
            ["7. 바둑미션과 문제은행 운영", "15. 장애 대응과 정기 체크리스트"],
            ["8. 게시판 기관 상담 1대1 문의 운영", ""],
        ],
        widths=[3.2, 3.2],
        font_size=8.8,
    )
    doc.add_page_break()

    # 1
    m.major("1 서비스 현황과 운영 전 주의사항")
    m.add_body("다음 표는 2026년 9월 10일 기준으로 공개 설정, API 응답, 강의 화면과 CI를 확인한 결과입니다. 설정이나 배포 버전이 바뀌면 이 표와 문서 버전도 함께 갱신합니다.")
    m.add_table(
        ["항목", "현재 상태", "운영 판단"],
        [
            ["홈페이지와 API", "공개 중 API readiness 정상", "기본 화면과 API 사용 가능"],
            ["토스페이먼츠", "테스트 모드 문서용 테스트 키", "실제 청구 금지 테스트만 가능"],
            ["소셜 로그인", "네이버 카카오 Google 활성", "인증 시작 주소의 공급자 이동 확인 완료"],
            ["게시판 서버 연결", "boardApiEnabled true", "게시글을 서버에 저장하고 다른 환경에서 재확인"],
            ["기존 강의 페이지 연결", "lectureApiEnabled true", "공개 강의 API와 운영 CMS 사용"],
            ["개발 역할 선택기", "demoRoleSwitcher false", "운영 화면에서 숨김"],
            ["Resend 발신 도메인", "최근 확인에서 인증 완료", "SMTP 수신 시험과 워커 확인 필요"],
            ["운영 워커", "전체 운영 스택 기동 여부 재확인 필요", "메일 영상 파일 큐 상태 확인 필수"],
            ["CI 보안 점검", "npm audit 취약점 0건", "웹 API 컨테이너 공급망 작업 모두 통과"],
        ],
        widths=[1.65, 2.15, 2.6],
        font_size=8.8,
    )
    m.add_note("게시 확인", "게시글 작성 뒤 목록에 다시 나타나는지 확인하고 다른 브라우저나 비공개 창에서도 같은 글이 조회되는지 점검합니다. 저장 실패가 표시되면 같은 내용을 반복 등록하지 말고 API 상태와 요청 ID를 먼저 확인합니다.")
    m.add_note("결제 주의", "테스트 결제는 실제 매출이 아닙니다. 운영 결제로 전환하기 전 가맹점 전용 라이브 클라이언트 키와 라이브 시크릿 키를 같은 상점의 한 쌍으로 적용하고 성공 실패 취소 환불을 다시 검증합니다.", danger=True)

    doc.add_heading("운영에서 반드시 지킬 원칙", level=2)
    m.add_checklist([
        "콘텐츠는 임시저장 또는 비공개 상태에서 검수한 뒤 공개한다.",
        "공개 콘텐츠를 삭제하지 않고 보관 상태로 전환해 기존 학습 기록을 보존한다.",
        "결제 완료 상태를 데이터베이스에서 직접 바꾸지 않는다.",
        "환불 전에 주문번호 결제키 금액 사용자를 서로 대조한다.",
        "비밀번호 SMTP 암호 OAuth 시크릿 토스 시크릿을 문서 화면 캡처 Git에 남기지 않는다.",
    ])

    # 2
    m.major("2 운영 역할과 권한")
    m.add_body("공개 회원가입으로 만들 수 있는 계정과 운영 권한 계정을 구분합니다. 운영자와 관리자는 일반 회원가입 화면에서 선택할 수 없으며, 승인된 서버 관리 절차로만 부여합니다.")
    m.add_table(
        ["역할", "할 수 있는 일", "할 수 없는 일"],
        [
            ["학생", "강의 미션 진도 대시보드 보호자 초대", "관리자 화면과 타인 정보 조회"],
            ["보호자", "초대 수락 연결 학생 리포트 확인", "연결되지 않은 학생 조회"],
            ["지도자", "수업도우미 자료와 승인된 반 학생 관리", "기관 결제 환불과 전체 회원 조회"],
            ["기관 관리자", "활성 기관의 라이선스 좌석 환불 요청 범위", "외부 결제 취소 직접 실행"],
            ["운영자", "강의 미션 상담 문의 결제 대사 커뮤니티 운영", "운영자 권한 자체 발급"],
            ["관리자", "전체 운영 정책 권한과 감사 검토", "승인 기록 없는 임의 변경"],
        ],
        widths=[1.05, 3.15, 2.2],
        font_size=8.7,
    )
    doc.add_heading("권한 확인 방법", level=2)
    m.add_steps([
        "https://handol-edu.com/account 에서 로그인합니다.",
        "계정과 보안 화면에서 표시 이름 이메일 인증 상태 역할 배지를 확인합니다.",
        "운영자 메뉴 주소로 이동합니다. 권한 안내가 나오면 새 계정을 만들지 말고 관리자에게 역할 확인을 요청합니다.",
        "여러 사람이 하나의 운영 계정을 공유하지 않습니다. 담당자별 계정을 사용해야 변경 이력을 추적할 수 있습니다.",
    ])
    m.add_note("권한 변경", "운영자 관리자 기관 관리자 역할 변경은 서비스 화면의 일반 기능이 아닙니다. 승인자 변경 사유 적용자 적용 시각을 남긴 뒤 서버 담당자가 처리합니다.")

    doc.add_heading("주요 관리자 주소", level=2)
    m.add_table(
        ["업무", "주소", "권한"],
        [
            ["강의 콘텐츠", "/admin/lessons", "운영자 관리자"],
            ["바둑문제 입력기", "/admin/missions", "운영자 관리자"],
            ["결제 대사", "/admin/payments", "운영자 관리자"],
            ["기관 상담", "/admin/consultations", "운영자 관리자"],
            ["1대1 문의", "/admin/inquiries", "운영자 관리자"],
            ["커뮤니티 신고", "/admin/community-reports", "운영자 관리자"],
            ["운영 워커", "/admin/operations", "운영자 관리자"],
            ["기관 라이선스", "/organization/admin", "기관 관리자"],
        ],
        widths=[1.55, 2.8, 2.05],
    )
    m.add_body("모든 상대 주소 앞에는 https://handol-edu.com 을 붙입니다.", style="Small Note")

    # 3
    m.major("3 로그인 계정 보호자 연결")
    doc.add_heading("일반 회원 로그인과 회원가입", level=2)
    m.add_steps([
        "계정 화면에서 로그인 또는 회원가입 탭을 선택합니다.",
        "회원가입 때 학생 또는 보호자 역할과 연령대를 정확히 선택합니다.",
        "이메일 인증 안내를 받은 뒤 인증 링크를 열거나 인증 화면에 토큰을 입력합니다.",
        "만 14세 미만 학생은 보호자 동의가 완료될 때까지 제한 상태가 표시됩니다.",
        "로그인 후 표시 이름 이메일 인증 역할이 맞는지 확인합니다.",
    ])
    doc.add_heading("비밀번호 재설정", level=2)
    m.add_steps([
        "로그인 화면의 비밀번호를 잊으셨나요를 선택합니다.",
        "가입 이메일로 재설정 안내를 요청합니다.",
        "메일의 유효시간 안에 링크를 열고 새 비밀번호를 입력합니다.",
        "메일이 없으면 스팸함을 확인하고 11장의 이메일 점검 절차를 따릅니다.",
    ])
    doc.add_heading("보호자 연결", level=2)
    m.add_steps([
        "학생 계정으로 /guardian 에 접속해 보호자 이메일로 초대를 만듭니다.",
        "보호자 계정으로 로그인한 뒤 받은 초대 링크 또는 토큰을 확인합니다.",
        "보호자는 동의 내용을 읽고 동의한 뒤 초대를 수락합니다.",
        "연결 상태가 활성으로 표시되는지 확인합니다.",
        "연결이 불필요해지면 학생 또는 보호자 화면에서 연결 해제를 실행합니다.",
    ])
    m.add_note("계정 탈퇴", "계정과 보안의 탈퇴 안내를 읽고 연결된 소셜 계정으로 본인 확인한 뒤 탈퇴합니다. 학습 정보와 로그인 수단은 삭제 또는 익명화되지만 결제 환불 감사 기록은 법정 보존 기준에 따라 분리 보관될 수 있습니다.", danger=True)

    # 4
    m.major("4 매일 해야 하는 운영 업무")
    doc.add_heading("업무 시작 체크", level=2)
    m.add_checklist([
        "홈페이지 첫 화면과 로그인 화면이 정상적으로 열린다.",
        "API readiness가 정상이다.",
        "운영 워커 상태 화면의 전체 상태가 정상이다.",
        "결제 대사 화면의 확인 필요 건수를 확인했다.",
        "새 기관 상담과 1대1 문의를 확인했다.",
        "문의 답변 이메일 발송 실패가 있는지 확인했다.",
        "커뮤니티 미처리 신고를 확인했다.",
    ])
    doc.add_heading("권장 처리 순서", level=2)
    m.add_steps([
        "즉시 확인 상태의 워커 장애와 결제 불일치를 먼저 처리합니다.",
        "개인정보 노출 불법정보 등 긴급 신고를 확인하고 필요하면 게시글을 숨깁니다.",
        "1대1 문의와 기관 상담을 오래된 접수 순으로 처리합니다.",
        "예약된 콘텐츠와 오늘 공개할 강의 미션을 검수합니다.",
        "업무 종료 전에 처리 결과와 미처리 사유를 운영 일지에 기록합니다.",
    ])
    doc.add_heading("운영 일지에 남길 내용", level=2)
    m.add_table(
        ["항목", "기록 예"],
        [
            ["시각과 담당자", "2026-09-09 09:10 김운영"],
            ["처리 업무", "결제 대사 상담 문의 콘텐츠 공개"],
            ["대상 식별자", "주문번호 문의번호 강의 ID 미션 ID"],
            ["결과", "완료 보류 장애 전달"],
            ["근거", "화면 상태 요청 ID 공급자 조회 결과"],
            ["다음 조치", "담당자와 예정 시각"],
        ],
        widths=[1.5, 4.9],
    )
    m.add_note("기록 금지", "운영 일지에는 비밀번호 카드번호 계좌번호 토스 시크릿 OAuth 시크릿 SMTP 암호 전체 이메일 본문을 넣지 않습니다.", danger=True)

    # 5
    m.major("5 홈페이지와 시대별 학습 콘텐츠 운영")
    m.add_body("홈페이지의 시대 퀴즈와 바둑미션은 같은 시대 식별자로 연결됩니다. 사용자가 퀴즈를 맞히면 해당 시대에 등록되고 공개된 바둑미션이 열려야 하며, 미션을 완료한 뒤에는 홈페이지로 돌아옵니다.")
    doc.add_heading("시대별 콘텐츠 공개 순서", level=2)
    m.add_steps([
        "강의 CMS에서 시대와 강의 ID를 확인하고 새 강의를 비공개로 등록합니다.",
        "바둑문제 입력기에서 같은 시대와 연결 강의를 지정해 문제를 만듭니다.",
        "문제를 임시저장하고 자동검수와 기록 없는 미리보기를 완료합니다.",
        "문제를 게시합니다. 같은 시대에 여러 문제가 있으면 공개 순서와 가장 먼저 열릴 문제를 확인합니다.",
        "강의 영상과 자료가 준비되면 강의를 공개합니다.",
        "비로그인 창에서 홈페이지 시대 탭 퀴즈를 풀고 해당 시대 문제가 바로 열리는지 확인합니다.",
        "정답 완료 후 닫기 또는 완료 흐름이 홈페이지로 돌아오는지 확인합니다.",
    ])
    doc.add_heading("홈페이지에서 직접 바꾸기 어려운 항목", level=2)
    m.add_bullets([
        "시대 탭 문구 역사 이야기 퀴즈 질문과 정답은 현재 정적 홈페이지 코드에 포함될 수 있으므로 변경 시 배포가 필요합니다.",
        "강의와 바둑미션 데이터는 관리자 CMS에서 등록합니다.",
        "메인 교재 상품의 가격과 설명은 서버 상품 데이터와 화면 표시값이 일치해야 합니다.",
        "푸터의 사업자 정보 고객센터 전화번호 이용약관 링크는 실제 값으로 교체하고 법무 검토 후 공개합니다.",
    ])
    doc.add_heading("공개 후 확인", level=2)
    m.add_checklist([
        "PC와 모바일에서 시대 탭이 선택된다.",
        "정답과 오답 안내가 정확하다.",
        "정답 후 시대별 미션 주소에 eraId와 autostart 값이 전달된다.",
        "다른 시대의 전체 목록이 먼저 열리지 않는다.",
        "미션 완료 후 홈페이지로 돌아온다.",
        "비공개 또는 보관 콘텐츠가 일반 사용자에게 보이지 않는다.",
    ])

    # 6
    m.major("6 강의 콘텐츠 관리")
    m.add_url("관리 화면", "https://handol-edu.com/admin/lessons")
    m.add_body("강의는 비공개로 등록한 뒤 영상과 여섯 단계가 준비되었을 때 공개합니다. 운영 중에는 삭제 대신 보관을 사용합니다.")
    doc.add_heading("새 강의 등록", level=2)
    m.add_steps([
        "새 강의 등록을 누릅니다.",
        "강의 ID 시대 순서 단계 과정 강사 제목 요약 난이도 예상 시간을 입력합니다.",
        "누구나 볼 수 있는 체험 강의만 무료 샘플을 선택합니다.",
        "비공개 강의 등록을 눌러 저장합니다.",
        "왼쪽 전체 강의 목록에서 새 강의를 다시 선택해 입력값을 확인합니다.",
    ])
    m.add_table(
        ["입력 항목", "운영 기준"],
        [
            ["강의 ID", "중복되지 않는 영문 대문자와 숫자 조합 예 PRE-02"],
            ["시대와 순서", "홈페이지 학습 순서와 일치"],
            ["제목과 요약", "학부모와 학생이 이해할 수 있는 표현 사용"],
            ["예상 시간", "1분부터 600분 현재 화면 제한"],
            ["무료 샘플", "대표 체험 콘텐츠에만 적용"],
        ],
        widths=[1.65, 4.75],
    )
    doc.add_heading("영상 업로드와 검사", level=2)
    m.add_steps([
        "강의를 선택한 뒤 MP4 파일을 선택하고 영상 업로드를 누릅니다.",
        "업로드 중 검사 대기 검사 중 상태가 차례로 바뀌는지 기다립니다.",
        "검사 통과와 현재 영상 표시를 확인합니다.",
        "검사 오류이면 검사 다시 시도를 누르고 운영 워커 상태를 확인합니다.",
        "HLS 변환 오류이면 HLS 변환 다시 시도를 누릅니다.",
        "외부에서 준비한 HLS를 사용할 때만 master.m3u8 경로를 입력하고 준비된 HLS 연결을 사용합니다.",
    ])
    doc.add_heading("썸네일과 학습자료", level=2)
    m.add_bullets([
        "썸네일 허용 형식은 JPG JPEG PNG WebP입니다.",
        "학습자료는 PDF PPT PPTX DOC DOCX HWP HWPX를 선택할 수 있습니다.",
        "모든 파일은 격리와 형식 검사 악성코드 검사를 통과해야 활성화됩니다.",
        "파일이 거부되면 확장자만 바꾸지 말고 원본 프로그램에서 안전한 새 파일로 다시 저장합니다.",
    ])
    doc.add_heading("강의 상태 변경", level=2)
    m.add_table(
        ["상태", "사용 시점", "일반 사용자 노출"],
        [
            ["비공개", "작성 검수 수정 중", "안 됨"],
            ["공개", "영상 연결과 여섯 단계 완료 후", "접근권한에 따라 노출"],
            ["보관", "더 이상 신규 제공하지 않을 때", "안 됨 기존 기록 보존"],
        ],
        widths=[1.2, 3.35, 1.85],
    )
    m.add_note("공개 버튼 비활성", "영상이 연결되지 않았거나 기본 단계가 6개가 아니면 공개할 수 없습니다. 먼저 영상 검사 상태와 단계 수를 확인하세요.")

    # 7
    m.major("7 바둑미션과 문제은행 운영")
    m.add_url("관리 화면", "https://handol-edu.com/admin/missions")
    m.add_body("바둑미션은 초기 바둑판과 사용자가 둘 수 있는 정답 수순 트리를 데이터로 저장합니다. 한 문제를 만들 때 기본정보 초기 판 정답 수순 힌트 해설 점수를 모두 입력한 뒤 자동검수와 미리보기를 통과해야 합니다.")
    doc.add_heading("문제 등록 절차", level=2)
    m.add_steps([
        "새 문제를 누르고 시대 과정 권 강 문제군 카테고리 난이도 착수 색 문제 유형을 입력합니다.",
        "9줄 13줄 19줄 중 판 크기를 선택합니다.",
        "흑돌 백돌 지우기 도구로 문제 시작 시점의 바둑돌을 놓습니다.",
        "대표 정답을 입력하고 여러 수 문제라면 상대 자동 응수와 다음 사용자 착수를 수순 트리로 연결합니다.",
        "복수 정답 허용 변화 금지 수순이 있으면 각각 별도 가지로 등록합니다.",
        "힌트 해설 기본 점수 오답 감점 힌트 감점을 입력합니다.",
        "임시저장 후 자동검수를 실행합니다.",
        "기록 없는 미리보기에서 실제 사용자처럼 착수해 모든 정답과 오답 피드백을 확인합니다.",
        "시대와 연결 강의가 맞는지 다시 확인하고 게시합니다.",
    ])
    doc.add_heading("문제 공개 전 필수 확인", level=2)
    m.add_checklist([
        "초기 흑돌과 백돌이 같은 자리에 겹치지 않는다.",
        "착수 차례와 사용자 돌 색이 맞다.",
        "대표 정답의 좌표가 바둑판 범위 안에 있다.",
        "상대 응수 뒤 사용자의 다음 수가 이어진다.",
        "정답 가지가 끝나고 성공 상태로 판정된다.",
        "불법 착수 자충 패 이미 놓인 자리가 올바르게 거부된다.",
        "시대 ID와 강의 ID가 실제 공개 콘텐츠와 연결된다.",
        "힌트와 해설에 정답 좌표를 잘못 적지 않았다.",
        "점수와 보상이 중복 지급되지 않는다.",
    ])
    doc.add_heading("문제 상태 운영", level=2)
    m.add_table(
        ["상태", "의미", "운영 행동"],
        [
            ["임시저장", "작성 중", "자동검수와 미리보기 전 공개 금지"],
            ["검수 요청", "다른 담당자 확인 대기", "정답 수순과 설명 교차 검토"],
            ["게시", "사용자에게 제공", "학습 기록 연결 상태 유지"],
            ["보관", "신규 노출 중단", "삭제하지 않고 기존 기록 보존"],
        ],
        widths=[1.25, 2.25, 2.9],
    )
    m.add_note("상세 설명", "초기 바둑판과 정답 수순 트리 작성법은 docs 폴더의 바둑문제 입력기 운영 매뉴얼을 함께 사용하세요. 전체 운영 매뉴얼에서는 게시와 연결 절차를 중심으로 다룹니다.")

    # 8
    m.major("8 게시판 기관 상담 1대1 문의 운영")
    doc.add_heading("게시판 현재 운영 방식", level=2)
    m.add_body("공지사항 수업 팁 여행기 FAQ 기관상담 지도자 수업도우미 교재자료는 board.html의 게시판 유형으로 접근합니다. 현재 게시판 API가 활성화되어 작성한 글은 서버에 저장됩니다. 게시 직후 목록, 상세 화면, 다른 브라우저에서 같은 내용이 보이는지 확인해야 게시 완료로 판단합니다.")
    m.add_table(
        ["게시판", "주소"],
        [
            ["공지사항", "/board.html?type=notice"],
            ["수업 팁", "/board.html?type=classTip"],
            ["여행기", "/board.html?type=travel"],
            ["자주 묻는 질문", "/board.html?type=faq"],
            ["1대1 문의", "/board.html?type=inquiry"],
            ["기관상담", "/board.html?type=consultation"],
            ["지도자 수업도우미", "/board.html?type=classHelper"],
            ["교재자료", "/board.html?type=resource"],
        ],
        widths=[2.05, 4.35],
    )
    m.add_note("중요", "저장 버튼을 여러 번 누르면 중복 글이 생길 수 있습니다. 응답이 늦으면 목록을 새로 불러와 저장 여부를 먼저 확인하고, 첨부파일이나 개인정보가 포함된 글은 공개 범위를 다시 점검합니다.", danger=True)

    doc.add_heading("기관 도입 상담", level=2)
    m.add_url("관리 화면", "https://handol-edu.com/admin/consultations")
    m.add_steps([
        "상태 기관 유형 검색어로 접수 목록을 조회합니다.",
        "접수 항목을 선택해 기관명 담당자 연락처 이메일 예상 인원 동의 기록 문의 내용을 확인합니다.",
        "실제 검토를 시작하면 검토 중으로 변경합니다.",
        "전화나 이메일 연락이 끝나면 연락 완료로 변경합니다.",
        "더 이상 후속 조치가 없으면 종료로 변경합니다.",
        "재상담이 필요하면 종료 또는 연락 완료 상태를 검토 중으로 되돌립니다.",
    ])
    m.add_table(
        ["상태", "뜻"],
        [["접수", "아직 담당자가 확인하지 않음"], ["검토 중", "담당자가 확인하고 대응 중"], ["연락 완료", "기관 담당자에게 연락함"], ["종료", "상담 후속 조치가 끝남"]],
        widths=[1.45, 4.95],
    )

    doc.add_heading("1대1 문의", level=2)
    m.add_url("관리 화면", "https://handol-edu.com/admin/inquiries")
    m.add_steps([
        "상태와 문의 유형으로 목록을 좁히고 오래된 접수부터 선택합니다.",
        "문의 내용과 첨부파일을 확인합니다. 첨부파일을 외부에 재전송하지 않습니다.",
        "처리가 필요하면 검토 중으로 변경합니다.",
        "운영자 답변을 입력하고 답변 등록을 누릅니다.",
        "답변 이메일 발송 상태가 발송 완료인지 확인합니다.",
        "발송 실패이며 재시도 버튼이 표시되면 원인을 확인한 뒤 이메일 재시도를 누릅니다.",
        "완료된 문의는 종료로 변경합니다.",
    ])
    m.add_note("답변 수정", "답변 완료 문의를 검토 중으로 되돌리면 기존 답변이 삭제됩니다. 수정 전 기존 답변과 재검토 사유를 운영 일지에 남기세요.", danger=True)

    doc.add_heading("커뮤니티 신고", level=2)
    m.add_url("관리 화면", "https://handol-edu.com/admin/community-reports")
    m.add_bullets([
        "개인정보 노출 불법정보 신체 위협은 먼저 게시글 숨김을 검토합니다.",
        "게시글 숨김은 같은 게시글의 미처리 신고를 함께 종결하므로 원문과 사유를 먼저 확인합니다.",
        "위반이 아니면 신고 기각을 선택합니다.",
        "처리 시각과 판단 근거를 감사 기록에 남깁니다.",
    ])

    # 9
    m.major("9 구독 결제 대사 환불 운영")
    m.add_url("사용자 구독", "https://handol-edu.com/subscriptions")
    m.add_url("관리자 결제 대사", "https://handol-edu.com/admin/payments")
    m.add_body("구독 결제는 서버가 주문 금액을 만들고 토스페이먼츠 승인 결과를 다시 검증한 뒤 구독을 발급합니다. 브라우저에 표시된 금액만 보고 결제 완료로 판단하지 않습니다.")
    doc.add_heading("현재 결제 모드", level=2)
    m.add_bullets([
        "현재 공개 설정은 테스트 모드입니다.",
        "문서용 공용 테스트 키가 적용되어 있어 실제 금액은 청구되지 않습니다.",
        "실제 판매 전에 같은 상점에서 발급한 가맹점 전용 라이브 결제위젯 키와 라이브 시크릿 키로 교체해야 합니다.",
    ])
    doc.add_heading("테스트 결제 확인", level=2)
    m.add_steps([
        "테스트 회원으로 로그인하고 구독 화면을 엽니다.",
        "플랜을 선택해 주문을 만들고 결제수단과 약관 위젯이 표시되는지 확인합니다.",
        "토스 테스트 결제를 완료합니다.",
        "구독 내역과 주문 내역이 결제 완료로 바뀌는지 확인합니다.",
        "토스 개발자센터 테스트 결제 내역에서 같은 주문번호를 확인합니다.",
        "성공 화면을 새로고침해도 구독이 중복 발급되지 않는지 확인합니다.",
    ])
    doc.add_heading("매일 결제 대사", level=2)
    m.add_steps([
        "기본 최근 30일 조건으로 조회합니다.",
        "확인 필요 건수를 선택해 불일치 사유를 읽습니다.",
        "결제키가 있으면 토스 재조회 동기화를 실행합니다.",
        "결제키가 없을 때만 토스 관리자 화면에서 주문번호와 금액을 확인한 뒤 결제키를 직접 입력합니다.",
        "결과가 일치로 바뀌었는지 확인합니다.",
        "필요하면 현재 조회 조건의 CSV를 내려받아 승인된 보관소에 저장합니다.",
    ])
    doc.add_heading("전액 환불", level=2)
    m.add_steps([
        "주문번호 사용자 결제금액 토스 결제키 환불 가능 금액을 대조합니다.",
        "고객의 환불 요청과 내부 승인 근거를 확인합니다.",
        "전액 환불을 열고 5자 이상의 구체적인 환불 사유를 입력합니다.",
        "구독 환불은 즉시 권한 회수 확인란을 선택합니다.",
        "전액 환불 확정을 누른 뒤 주문이 취소 상태이고 환불 누적액이 맞는지 확인합니다.",
        "구독 권한이 회수되었는지 사용자 구독 내역과 대사 화면에서 확인합니다.",
    ])
    m.add_note("금지", "토스 관리자 화면과 서비스 DB가 다르다고 해서 DB 값을 직접 결제 완료 또는 취소로 수정하지 않습니다. 먼저 재조회 동기화를 사용하고, 그래도 해결되지 않으면 주문번호와 요청 ID를 서버 담당자에게 전달합니다.", danger=True)

    # 10
    m.major("10 학생 보호자 지도자 기관 기능")
    doc.add_heading("학생", level=2)
    m.add_bullets([
        "/lessons 에서 시대별 공개 강의를 확인합니다.",
        "/dashboard 에서 전체 진행률 시대별 완료율 최근 학습과 다음 강의를 확인합니다.",
        "/missions 에서 판 크기 과정 문제군 유형 난이도 풀이 상태로 문제를 찾습니다.",
        "로그인 사용자는 즐겨찾기와 이어하기 학습 기록을 사용할 수 있습니다.",
    ])
    doc.add_heading("보호자", level=2)
    m.add_bullets([
        "/guardian 에서 연결된 학생을 선택해 학습 리포트를 확인합니다.",
        "활성 연결과 동의 범위가 있는 학생만 조회됩니다.",
        "잘못 연결된 학생은 즉시 연결 해제하고 고객지원에 기록을 남깁니다.",
    ])
    doc.add_heading("지도자", level=2)
    m.add_bullets([
        "지도자 수업도우미는 한 게시물 안에서 영상 PPT 활동지 퀴즈 미션 해설 진행 가이드를 함께 여는 구조입니다.",
        "반과 학생 정보는 담당 범위 안에서만 확인합니다.",
        "기관 결제와 환불은 기관 관리자 또는 운영자에게 요청합니다.",
    ])
    doc.add_heading("기관 관리자", level=2)
    m.add_url("기관 관리", "https://handol-edu.com/organization/admin")
    m.add_bullets([
        "활성 기관 관리자 멤버십이 있는 기관만 표시됩니다.",
        "기관 라이선스 좌석과 환불 요청 범위를 확인합니다.",
        "토스 결제 취소 실행은 운영자 또는 관리자에게 요청합니다.",
    ])
    doc.add_heading("교재 QR", level=2)
    m.add_steps([
        "QR 주소가 /qr/코드 형식인지 확인합니다.",
        "로그인이 필요한 경우 로그인 후 원래 QR 강의로 돌아오는지 확인합니다.",
        "코드가 만료되었거나 존재하지 않으면 새 코드를 임의 생성하지 말고 교재 식별정보와 함께 운영자에게 전달합니다.",
    ])

    # 11
    m.major("11 알림 이메일 소셜 로그인 운영")
    doc.add_heading("앱 알림", level=2)
    m.add_url("알림함", "https://handol-edu.com/notifications")
    m.add_bullets([
        "문의 답변과 서비스 안내가 계정별로 표시됩니다.",
        "개별 읽음 또는 모두 읽음을 사용할 수 있습니다.",
        "문의 답변 알림은 문의 상세로 이동하는 링크를 제공합니다.",
    ])

    doc.add_heading("Resend SMTP 기본값", level=2)
    m.add_table(
        ["설정", "운영 값 또는 기준"],
        [
            ["SMTP 호스트", "smtp.resend.com"],
            ["포트", "587"],
            ["보안", "SMTP_SECURE false 와 SMTP_REQUIRE_TLS true"],
            ["사용자", "resend"],
            ["발신자", "바둑타고 <no-reply@handol-edu.com>"],
            ["DKIM 선택자", "resend"],
            ["암호", "Resend가 발급한 값 서버 Secret에만 저장"],
        ],
        widths=[1.6, 4.8],
    )
    m.add_note("DNS", "Resend 도메인 화면의 DKIM과 SPF 레코드는 DNSZi에 이름과 값을 정확히 등록합니다. DKIM 값은 중간 공백 없이 한 줄로 저장하고, 공급자가 값을 바꾸면 Resend 화면의 최신 값을 사용합니다.")
    doc.add_heading("메일이 오지 않을 때", level=2)
    m.add_steps([
        "받는 주소와 스팸함 수신 거부 목록을 확인합니다.",
        "Resend Domains에서 handol-edu.com과 각 DNS 레코드가 Verified인지 확인합니다.",
        "운영 워커 상태에서 계정 인증 복구 메일과 문의 답변 알림 큐를 확인합니다.",
        "관리자 문의 화면에서 발송 대기 발송 중 발송 완료 발송 제외 발송 실패 상태를 확인합니다.",
        "발송 실패는 오류 원인을 해결한 뒤 화면의 이메일 재시도를 사용합니다.",
        "서버 로그에는 수신자 주소나 SMTP 암호를 복사하지 말고 오류 코드와 요청 ID만 기록합니다.",
    ])

    doc.add_heading("소셜 로그인", level=2)
    m.add_body("서버에는 공급자별 Client ID Client Secret Redirect URI가 모두 있어야 하고, 브라우저 config.js에서도 OAuth 사용과 제공자 목록을 켜야 로그인 버튼이 실제로 동작합니다.")
    m.add_table(
        ["공급자", "운영 Redirect URI"],
        [
            ["네이버", "https://handol-edu.com/api/v1/auth/oauth/naver/callback"],
            ["카카오", "https://handol-edu.com/api/v1/auth/oauth/kakao/callback"],
            ["Google", "https://handol-edu.com/api/v1/auth/oauth/google/callback"],
        ],
        widths=[1.35, 5.05],
        font_size=8.6,
    )
    m.add_steps([
        "각 공급자 콘솔에 운영 도메인과 Redirect URI를 정확히 등록합니다.",
        "필요한 동의 항목은 네이버와 Google의 기본 프로필 이메일, 카카오의 닉네임과 이메일만 설정합니다.",
        "config.js에서 oauthEnabled를 true로 하고 oauthProviders에 naver kakao google을 등록합니다.",
        "배포 후 신규 로그인 기존 계정 연결 연결 해제 탈퇴 본인 확인을 공급자별로 테스트합니다.",
        "오류 시 공급자 콘솔 값과 서버 환경변수의 대소문자 경로 슬래시까지 비교합니다.",
    ])
    m.add_note("현재 상태", "2026년 9월 10일 기준 공개 config.js에서 네이버, 카카오, Google 로그인이 활성화되어 있으며 세 공급자의 인증 화면으로 이동하는 HTTP 302 응답을 확인했습니다. 실제 계정 로그인과 최초 동의 화면은 공급자별 테스트 계정으로 정기 점검합니다.")

    # 12
    m.major("12 서버 상태 점검과 로그 확인")
    m.add_body("다음 명령은 서버 담당자가 SSH로 운영 서버에 로그인한 뒤 /opt/hanstone에서 실행하는 기준입니다. 실제 배포 경로가 다르면 승인된 경로를 사용합니다.")
    doc.add_heading("읽기 전용 기본 점검", level=2)
    m.add_code([
        "cd /opt/hanstone",
        "docker compose ps",
        "curl -fsS https://handol-edu.com/api/v1/health/ready",
        "docker compose logs --since=30m api | tail -n 200",
    ])
    doc.add_heading("공개 운영 설정 확인", level=2)
    m.add_code([
        "curl -fsS -H 'Cache-Control: no-cache' 'https://handol-edu.com/config.js?verify=현재시각'",
    ])
    m.add_table(
        ["설정", "운영 기준"],
        [
            ["oauthEnabled", "true"],
            ["oauthProviders", "naver kakao google"],
            ["boardApiEnabled", "true"],
            ["lectureApiEnabled", "true"],
            ["demoRoleSwitcher", "false"],
            ["tossPayments mode", "test 실결제 전까지 유지"],
        ],
        widths=[2.3, 4.1],
        font_size=8.8,
    )
    m.add_table(
        ["결과", "판정"],
        [
            ["API status ok database ok", "기본 API와 DB 연결 정상"],
            ["컨테이너 healthy", "상태검사 통과"],
            ["unhealthy 또는 restarting", "로그 확인 후 배포 담당자 호출"],
            ["ECONNREFUSED 127.0.0.1:587", "컨테이너가 최신 SMTP 환경값 없이 시작됐을 가능성"],
            ["401 또는 403", "로그인 세션 또는 역할 부족"],
            ["요청 ID가 포함된 5xx", "요청 ID와 시각으로 서버 로그 대조"],
        ],
        widths=[2.6, 3.8],
        font_size=8.8,
    )
    doc.add_heading("운영 워커", level=2)
    m.add_bullets([
        "account-mail-worker는 이메일 인증과 비밀번호 복구 메일을 처리합니다.",
        "inquiry-notification-worker는 1대1 문의 답변 이메일을 처리합니다.",
        "video-scan-worker는 업로드 영상을 악성코드 검사합니다.",
        "hls-transcode-worker는 영상을 스트리밍 HLS로 변환합니다.",
        "video-cleanup-worker는 교체되거나 연결되지 않은 파일을 보존 정책에 따라 정리합니다.",
    ])
    m.add_url("화면 점검", "https://handol-edu.com/admin/operations")
    m.add_body("전체 상태가 확인 필요이면 처리 대기와 최종 실패를 확인하고, 즉시 확인이면 오래된 잠금이나 장시간 적체가 있으므로 해당 워커 로그를 바로 확인합니다.")
    m.add_code([
        "docker compose --env-file deploy/production.env -f deploy/compose.production.yaml ps",
        "docker compose --env-file deploy/production.env -f deploy/compose.production.yaml logs --since=30m account-mail-worker inquiry-notification-worker video-scan-worker hls-transcode-worker video-cleanup-worker",
    ])
    m.add_note("명령 복사", "환경 파일 내용을 cat grep printenv로 출력하지 않습니다. 설정 여부만 확인하는 검증 스크립트를 사용하고 화면 공유 전에 Secret 노출 여부를 확인합니다.", danger=True)

    # 13
    m.major("13 배포 변경 승인 롤백")
    m.add_body("배포는 코드 파일을 서버에서 직접 고치는 작업이 아닙니다. 검증된 커밋과 배포 묶음을 준비하고 사전 점검을 통과한 뒤 원자적으로 전환합니다.")
    doc.add_heading("배포 전 승인", level=2)
    m.add_checklist([
        "변경 내용과 영향 화면이 기록되었다.",
        "웹과 API 자동 테스트가 통과했다.",
        "DB 마이그레이션 영향과 되돌림 전략이 검토되었다.",
        "OAuth 메일 결제 파일 저장소 설정이 Secret으로 주입되었다.",
        "운영 프리플라이트가 모두 통과했다.",
        "직전 정상 버전과 롤백 담당자가 확인되었다.",
        "배포 시간과 사용자 공지가 승인되었다.",
    ])
    doc.add_heading("운영 배포 순서", level=2)
    m.add_steps([
        "동일 커밋으로 웹 배포 묶음과 API 이미지를 준비합니다.",
        "API 이미지는 변경되지 않는 repository at sha256 digest로 기록합니다.",
        "DB 마이그레이션을 한 번만 적용합니다.",
        "운영 환경으로 production preflight를 실행하고 종료 코드 0과 모든 pass를 확인합니다.",
        "API와 다섯 워커를 같은 이미지로 교체합니다.",
        "정적 웹 릴리스를 전환합니다.",
        "배포 식별값 liveness readiness HTTPS 실제 화면을 검증합니다.",
        "결제 로그인 메일 강의 미션 상담의 핵심 흐름을 점검합니다.",
        "모든 결과를 변경승인 기록에 남기고 배포를 종료합니다.",
    ])
    m.add_code([
        "docker compose --env-file deploy/production.env -f deploy/compose.production.yaml run --rm api node dist/production-preflight.js",
        "docker compose --env-file deploy/production.env -f deploy/compose.production.yaml up -d api account-mail-worker inquiry-notification-worker video-scan-worker hls-transcode-worker video-cleanup-worker",
    ])
    doc.add_heading("정적 홈페이지 확인", level=2)
    m.add_code([
        "readlink -f /var/www/hanstone/current",
        "sudo python3 deploy/verify-hosting-release.py --base-url https://handol-edu.com --expected-commit 현재_40자리_커밋",
    ])
    m.add_note("롤백 판단", "검증 결과가 ok false 또는 rollbackRecommended true이면 정상 배포로 확정하지 않습니다. 직전 검증된 정적 릴리스와 불변 API 이미지로 되돌리되 적용된 DB 마이그레이션을 임의로 삭제하지 않습니다.", danger=True)
    m.add_note("상세 배포 절차", "서버 설치와 원자적 릴리스 전환은 deploy/HOSTING_INSTALL.md, 검증은 deploy/HOSTING_VERIFY.md, 변경 기록은 docs/RELEASE_CHANGE_APPROVAL_TEMPLATE.md를 따릅니다.")

    doc.add_heading("CI 보안 점검", level=2)
    m.add_body("main 브랜치에 푸시한 변경은 웹, API, API 컨테이너, 브라우저, 공급망 작업을 모두 통과해야 배포 후보로 사용합니다. 2026년 9월 10일 기준 multer 2.3.0과 nodemailer 10.0.2를 적용했으며 웹과 API의 운영 의존성 감사 결과는 취약점 0건입니다.")
    m.add_steps([
        "GitHub Actions에서 대상 커밋의 CI 실행을 엽니다.",
        "web, api, api-container, browser, supply-chain 작업이 모두 success인지 확인합니다.",
        "공급망 작업의 웹과 API npm audit가 모두 취약점 0건인지 확인합니다.",
        "실패한 작업이 있으면 로그에서 패키지 이름, 심각도, 고정 버전을 기록하고 lock 파일까지 함께 갱신합니다.",
        "수정 뒤 전체 테스트와 빌드를 다시 실행하고 새 CI가 성공한 커밋만 배포합니다.",
    ])
    m.add_code([
        "npm run ci",
        "node server/scripts/npm-audit-with-retry.mjs .",
        "node server/scripts/npm-audit-with-retry.mjs server",
    ])
    m.add_url("최근 검증된 CI", "https://github.com/gogumaz/HanStoneHomepage/actions/runs/34413191522")
    m.add_note("운영 적용 범위", "소스와 CI에서 보안 의존성을 수정했더라도 실행 중인 API 컨테이너는 자동으로 바뀌지 않을 수 있습니다. 정식 API 릴리스에서 새 이미지를 배포하고 readiness와 주요 기능을 다시 확인합니다.", danger=True)

    # 14
    m.major("14 보안 개인정보 백업 관리")
    doc.add_heading("비밀정보 관리", level=2)
    m.add_checklist([
        "시크릿 키는 서버 환경 파일 또는 비밀 관리 서비스에만 저장한다.",
        "config.js에는 공개 OAuth 제공자 목록과 토스 클라이언트 키만 둔다.",
        "토스 시크릿 OAuth Client Secret SMTP 암호 개인키는 브라우저 파일에 넣지 않는다.",
        "채팅 화면 캡처 저장소 로그에 노출된 키는 즉시 폐기하고 재발급한다.",
        "운영자 계정은 개인별로 사용하고 퇴사 또는 역할 변경 즉시 회수한다.",
    ])
    doc.add_heading("개인정보 처리", level=2)
    m.add_bullets([
        "기관 상담과 1대1 문의 연락처 내용 첨부파일은 업무에 필요한 담당자만 확인합니다.",
        "사용자 요청에 따라 열람 정정 삭제 동의 철회를 처리하고 처리 이력을 남깁니다.",
        "탈퇴 후 일반 계정정보와 학습 진도는 삭제 또는 익명화하고 법정 보존 결제 기록은 분리합니다.",
        "운영 로그에는 이메일 전화번호 상담 내용 결제수단 세부정보를 넣지 않습니다.",
    ])
    m.add_table(
        ["기록", "기준 보존기간"],
        [
            ["계약과 청약철회", "5년"],
            ["결제와 서비스 공급", "5년"],
            ["소비자 불만과 분쟁", "3년"],
            ["표시와 광고", "6개월"],
            ["일반 운영 로그", "기본 90일"],
            ["원시 분석 이벤트", "13개월"],
            ["비식별 집계 통계", "3년"],
        ],
        widths=[2.9, 3.5],
    )
    m.add_note("법무 확인", "보존기간과 파기 절차는 실제 사업자 정보 이용약관 개인정보처리방침 결제 정책을 법무 검토한 뒤 확정합니다.")

    doc.add_heading("백업과 복구", level=2)
    m.add_checklist([
        "PostgreSQL 시점 복구 PITR가 활성화되어 있다.",
        "백업 보존기간이 30일 이상이다.",
        "객체 저장소 버전 관리가 활성화되어 있다.",
        "목표 복구 시점 RPO 15분과 복구 시간 RTO 4시간 이하를 충족한다.",
        "최근 100일 이내 격리된 스테이징 환경에서 복구훈련을 완료했다.",
        "복구훈련 결과를 운영 데이터나 접속 문자열 없이 증빙으로 보관했다.",
    ])

    # 15
    m.major("15 장애 대응과 정기 체크리스트")
    doc.add_heading("장애 대응 기본 순서", level=2)
    m.add_steps([
        "장애 범위를 확인합니다. 한 사용자 한 기능 전체 사이트 중 어디까지인지 구분합니다.",
        "발생 시각 주소 사용 계정 역할 동작 순서 오류 문구 요청 ID를 기록합니다.",
        "API readiness 컨테이너 상태 운영 워커 상태를 읽기 전용으로 확인합니다.",
        "최근 배포나 설정 변경과 발생 시각을 비교합니다.",
        "개인정보와 Secret을 가린 뒤 담당자에게 전달합니다.",
        "결제 메일 데이터 손상 위험이 있으면 기능을 일시 중단하고 임의 재처리를 하지 않습니다.",
        "복구 뒤 같은 동작과 인접 기능을 재검증하고 원인 조치 재발 방지를 기록합니다.",
    ])
    doc.add_page_break()
    doc.add_heading("자주 발생하는 문제", level=2)
    m.add_table(
        ["증상", "먼저 확인", "조치"],
        [
            ["관리자 화면 권한 없음", "계정 역할과 세션", "다시 로그인 후 승인된 역할 요청"],
            ["인증메일 미수신", "Resend DNS SMTP 워커 발송 상태", "원인 해결 후 재시도"],
            ["결제 대기에서 멈춤", "토스 내역 주문번호 결제키", "관리자 대사에서 재조회 동기화"],
            ["결제창이 안 열림", "테스트와 라이브 키 종류 위젯 로딩", "같은 상점의 키 쌍과 config 확인"],
            ["게시글이 다른 PC에서 안 보임", "boardApiEnabled와 저장 API 응답", "요청 ID 확인 후 중복 등록 없이 담당자 호출"],
            ["강의 공개 버튼 비활성", "영상 연결과 단계 6개", "검사 완료와 단계 구성을 확인"],
            ["영상 검사 멈춤", "video scan 워커와 오래된 잠금", "워커 복구 후 화면 재시도"],
            ["퀴즈 뒤 전체 미션 목록", "시대 미션 게시와 eraId", "연결 데이터와 공개 순서 점검"],
            ["미션 완료 뒤 목록 복귀", "최신 웹 릴리스와 진입 경로", "캐시 새로고침 후 배포 버전 점검"],
            ["500 오류", "readiness와 요청 ID", "API 로그에서 같은 시각 요청 확인"],
        ],
        widths=[1.75, 2.3, 2.35],
        font_size=8.3,
    )

    doc.add_heading("매일", level=2)
    m.add_checklist([
        "홈페이지 로그인 API readiness 정상",
        "워커 전체 상태 정상",
        "결제 확인 필요 건 처리",
        "기관 상담 1대1 문의 긴급 신고 확인",
        "답변 이메일 발송 실패 확인",
    ])
    doc.add_heading("매주", level=2)
    m.add_checklist([
        "공개 강의와 미션 표본 재생 풀이 확인",
        "결제 대사 CSV 보관과 환불 이력 검토",
        "보관 또는 만료 예정 콘텐츠 확인",
        "사용하지 않는 운영 계정과 권한 확인",
        "백업 성공 기록과 저장공간 확인",
    ])
    doc.add_heading("매월", level=2)
    m.add_checklist([
        "OAuth 공급자 콘솔 오류와 Redirect URI 확인",
        "Resend 발신 평판 반송률 DNS 인증 확인",
        "개인정보 파기 대상과 법정 보존 분리 확인",
        "복구 가능한 백업 표본 확인",
        "SSL 인증서 만료일까지 30일 이상 남았는지 확인",
        "운영 매뉴얼과 실제 화면 차이를 반영해 버전 갱신",
    ])
    doc.add_heading("분기별", level=2)
    m.add_checklist([
        "격리된 스테이징 DB 복구훈련",
        "결제 성공 실패 취소 환불 중복 승인 재검증",
        "권한과 개인정보 접근 기록 검토",
        "장애 대응과 롤백 훈련",
        "콘텐츠 품질과 시대별 미션 누락 점검",
    ])

    # Appendix
    m.major("부록 화면 주소 빠른 찾기")
    m.add_table(
        ["화면", "전체 주소"],
        [
            ["홈페이지", "https://handol-edu.com/"],
            ["계정과 보안", "https://handol-edu.com/account"],
            ["강의 여행", "https://handol-edu.com/lessons"],
            ["바둑미션", "https://handol-edu.com/missions"],
            ["나의 여행지도", "https://handol-edu.com/dashboard"],
            ["보호자 연결", "https://handol-edu.com/guardian"],
            ["구독", "https://handol-edu.com/subscriptions"],
            ["알림함", "https://handol-edu.com/notifications"],
            ["강의 CMS", "https://handol-edu.com/admin/lessons"],
            ["바둑문제 입력기", "https://handol-edu.com/admin/missions"],
            ["결제 대사", "https://handol-edu.com/admin/payments"],
            ["기관 상담", "https://handol-edu.com/admin/consultations"],
            ["1대1 문의", "https://handol-edu.com/admin/inquiries"],
            ["커뮤니티 신고", "https://handol-edu.com/admin/community-reports"],
            ["운영 워커", "https://handol-edu.com/admin/operations"],
            ["기관 관리자", "https://handol-edu.com/organization/admin"],
            ["개인정보 안내", "https://handol-edu.com/privacy"],
            ["API 준비 상태", "https://handol-edu.com/api/v1/health/ready"],
        ],
        widths=[1.8, 4.6],
        font_size=8.5,
    )
    doc.add_heading("운영 인수인계 기록", level=2)
    m.add_table(
        ["확인 항목", "기록"],
        [
            ["인계 날짜와 시간", ""],
            ["인계자와 인수자", ""],
            ["현재 배포 커밋", ""],
            ["API와 워커 상태", ""],
            ["미처리 결제 상담 문의", ""],
            ["예약 공개 콘텐츠", ""],
            ["진행 중 장애", ""],
            ["다음 점검 예정", ""],
        ],
        widths=[2.0, 4.4],
    )
    m.add_note("문서 갱신", "화면 메뉴 상태값 권한 결제 메일 배포 절차가 바뀌면 코드 변경과 같은 작업 단위에서 이 운영 매뉴얼도 갱신합니다.")

    # Footer metadata and core properties.
    props = doc.core_properties
    props.title = "바둑타고 한국사 여행 홈페이지 전체 운영 매뉴얼"
    props.subject = "홈페이지 콘텐츠 고객지원 결제 메일 소셜 로그인 서버 운영 절차"
    props.author = "한돌에듀"
    props.keywords = "운영 매뉴얼, 홈페이지, 강의 CMS, 바둑미션, 결제, 상담, 서버"
    props.comments = "운영 기준일 2026-09-10"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    # Normalize the OOXML package so Word can render it reliably in automation.
    Document(OUTPUT).save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_manual()
